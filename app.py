from flask import Flask, render_template, request, jsonify
from flask_limiter import Limiter
from flask_limiter.util import get_remote_address
import os

from langchain_community.document_loaders import PyPDFLoader, Docx2txtLoader, CSVLoader
from langchain_text_splitters import RecursiveCharacterTextSplitter
from langchain_community.vectorstores import FAISS
from langchain_community.embeddings import HuggingFaceEmbeddings
from langchain_community.llms import Ollama

from docx import Document
from langchain_core.documents import Document as LCDocument

app = Flask(__name__)

# ── OWASP Fix 2: Rate limiting ─────────────────────────────────────────────
limiter = Limiter(
    get_remote_address,
    app=app,
    default_limits=["100 per hour"]
)
# ───────────────────────────────────────────────────────────────────────────

UPLOAD_FOLDER = "uploads"
VECTORSTORE_PATH = "vectorstore"

os.makedirs(UPLOAD_FOLDER, exist_ok=True)

embeddings = HuggingFaceEmbeddings()

if os.path.exists(VECTORSTORE_PATH):
    print("🔁 Loading existing vectorstore...")
    vector_db = FAISS.load_local(
        VECTORSTORE_PATH,
        embeddings,
        allow_dangerous_deserialization=True
    )
else:
    print("⚠ No vectorstore found. Upload a document first.")
    vector_db = None


# ---------------------------
# Document Loader
# ---------------------------
def load_document(path):
    try:
        path = path.lower()

        if path.endswith(".pdf"):
            loader = PyPDFLoader(path)
            return loader.load()

        elif path.endswith(".docx"):
            doc = Document(path)
            text = "\n".join([p.text for p in doc.paragraphs])
            return [LCDocument(page_content=text)]

        elif path.endswith(".csv"):
            loader = CSVLoader(path)
            return loader.load()

        else:
            print("❌ Unsupported extension:", path)
            return None

    except Exception as e:
        print("❌ Document loading error:", e)
        return None


# ---------------------------
# Routes
# ---------------------------
@app.route("/")
def index():
    return render_template("index.html")


# ---------------------------
# Upload Route
# ---------------------------
@app.route("/upload", methods=["POST"])
@limiter.limit("20 per hour")          # OWASP: rate limit uploads
def upload():
    global vector_db

    print("📤 Upload request received")

    if "file" not in request.files:
        return jsonify({"message": "No file part in request"}), 400

    file = request.files["file"]

    if file.filename == "":
        return jsonify({"message": "No file selected"}), 400

    path = os.path.join(UPLOAD_FOLDER, file.filename)
    file.save(path)
    print(f"✅ File saved: {path}")

    docs = load_document(path)
    if docs is None:
        return jsonify({"message": "Unsupported file type"}), 400

    print("📄 Document loaded")

    splitter = RecursiveCharacterTextSplitter(
        chunk_size=500,
        chunk_overlap=50
    )
    chunks = splitter.split_documents(docs)
    print(f"🧩 Created {len(chunks)} chunks")

    vector_db = FAISS.from_documents(chunks, embeddings)
    vector_db.save_local(VECTORSTORE_PATH)
    print("💾 Vectorstore saved")

    return jsonify({"message": "File uploaded and indexed successfully!"})


# ---------------------------
# Chat Route
# ---------------------------
@app.route("/chat", methods=["POST"])
@limiter.limit("30 per minute")        # OWASP: rate limit chat queries
def chat():
    global vector_db

    if vector_db is None and os.path.exists(VECTORSTORE_PATH):
        print("♻ Reloading vectorstore...")
        vector_db = FAISS.load_local(
            VECTORSTORE_PATH,
            embeddings,
            allow_dangerous_deserialization=True
        )

    if vector_db is None:
        return jsonify({"answer": "Please upload a document first."})

    data = request.json
    query = data.get("query", "")
    model = data.get("model", "llama3")

    print(f"💬 Query: {query}")
    print(f"🤖 Model: {model}")

    llm = Ollama(model=model)

    docs = vector_db.similarity_search(query, k=3)
    context = "\n".join([d.page_content for d in docs])

    prompt = f"""
You are a helpful assistant. Answer ONLY from the document context.

Context:
{context}

Question: {query}
"""

    try:
        response = llm.invoke(prompt)
        print("✅ Model replied")
    except Exception as e:
        print("❌ Ollama error:", e)
        response = "⚠ Model not responding. Try llama3:8b or gemma:2b."

    return jsonify({"answer": response})


# ---------------------------
# Reset Route
# ---------------------------
@app.route("/reset", methods=["POST"])
def reset():
    global vector_db

    vector_db = None

    if os.path.exists(VECTORSTORE_PATH):
        import shutil
        shutil.rmtree(VECTORSTORE_PATH)
        print("🗑 Vectorstore deleted")

    return jsonify({"status": "reset"})


# ---------------------------
# Run Server
# ---------------------------
if __name__ == "__main__":
    app.run(debug=False)               # OWASP Fix 1: debug off in production